from __future__ import annotations

import io
import uuid
from pathlib import Path
from typing import Any


from src.documents.models import Document


class DocumentParser:
    """
    Converts uploaded files into a normalized
    Document object.

    Supported formats:

    - TXT
    - PDF
    - DOCX
    """

    SUPPORTED_EXTENSIONS = {
        ".txt",
        ".pdf",
        ".docx",
    }

    def parse(
        self,
        uploaded_file,
    ) -> Document:
        """
        Parse a Streamlit UploadedFile or
        file-like object.
        """

        if uploaded_file is None:
            raise ValueError(
                "No document was provided."
            )

        filename = getattr(
            uploaded_file,
            "name",
            "uploaded_document",
        )

        extension = (
            Path(filename)
            .suffix
            .lower()
        )

        if (
            extension
            not in self.SUPPORTED_EXTENSIONS
        ):
            raise ValueError(
                "Unsupported file type: "
                f"{extension}. "
                "Supported formats are TXT, PDF and DOCX."
            )

        content_type = (
            self._get_content_type(
                extension
            )
        )

        if extension == ".txt":
            text = self._parse_txt(
                uploaded_file
            )

        elif extension == ".pdf":
            text = self._parse_pdf(
                uploaded_file
            )

        elif extension == ".docx":
            text = self._parse_docx(
                uploaded_file
            )

        else:
            raise ValueError(
                "Unsupported document format."
            )

        text = self._clean_text(
            text
        )

        if not text:
            raise ValueError(
                "The uploaded document "
                "does not contain readable text."
            )

        document_id = str(
            uuid.uuid4()
        )

        metadata = {
            "filename": filename,
            "content_type": content_type,
            "extension": extension,
            "character_count": len(text),
            "word_count": len(
                text.split()
            ),
        }

        return Document(
            document_id=document_id,
            filename=filename,
            content=text,
            content_type=content_type,
            metadata=metadata,
        )

    # =========================================================
    # TXT
    # =========================================================

    def _parse_txt(
        self,
        uploaded_file,
    ) -> str:

        raw = uploaded_file.read()

        if isinstance(
            raw,
            str,
        ):
            return raw

        for encoding in (
            "utf-8",
            "utf-8-sig",
            "latin-1",
        ):

            try:
                return raw.decode(
                    encoding
                )

            except UnicodeDecodeError:
                continue

        raise ValueError(
            "Unable to decode the text file."
        )

    # =========================================================
    # PDF
    # =========================================================

    def _parse_pdf(
        self,
        uploaded_file,
    ) -> str:

        try:
            import fitz

        except ImportError as exc:
            raise RuntimeError(
                "PyMuPDF is required for PDF parsing. "
                "Install it with: pip install PyMuPDF"
            ) from exc

        raw = uploaded_file.read()

        if not raw:
            return ""

        text_parts = []

        try:

            with fitz.open(
                stream=raw,
                filetype="pdf",
            ) as pdf:

                for page_number, page in enumerate(
                    pdf
                ):

                    page_text = page.get_text(
                        "text"
                    )

                    if page_text:
                        text_parts.append(
                            page_text
                        )

        except Exception as exc:

            raise RuntimeError(
                "Failed to parse PDF document."
            ) from exc

        return "\n\n".join(
            text_parts
        )

    # =========================================================
    # DOCX
    # =========================================================

    def _parse_docx(
        self,
        uploaded_file,
    ) -> str:

        try:
            from docx import Document as DocxDocument

        except ImportError as exc:
            raise RuntimeError(
                "python-docx is required for DOCX parsing. "
                "Install it with: pip install python-docx"
            ) from exc

        raw = uploaded_file.read()

        if not raw:
            return ""

        try:

            document = DocxDocument(
                io.BytesIO(raw)
            )

        except Exception as exc:

            raise RuntimeError(
                "Failed to open DOCX document."
            ) from exc

        paragraphs = []

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:
                paragraphs.append(
                    text
                )

        return "\n\n".join(
            paragraphs
        )

    # =========================================================
    # CLEANING
    # =========================================================

    def _clean_text(
        self,
        text: str,
    ) -> str:

        if not text:
            return ""

        lines = []

        for line in text.splitlines():

            cleaned = " ".join(
                line.split()
            )

            if cleaned:
                lines.append(
                    cleaned
                )

        return "\n".join(
            lines
        )

    # =========================================================
    # CONTENT TYPE
    # =========================================================

    def _get_content_type(
        self,
        extension: str,
    ) -> str:

        content_types = {
            ".txt": "text/plain",
            ".pdf": "application/pdf",
            ".docx": (
                "application/"
                "vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
        }

        return content_types.get(
            extension,
            "application/octet-stream",
        )